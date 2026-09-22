"""Source Data workbook: one sheet per figure panel, from the frozen figure data (Nature Portfolio policy).

    python c122_source_data.py        # -> source_data/SourceData.xlsx
"""
import json
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
FD = json.load(open(os.path.join(HERE, 'data', 'selfcheck', 'figure_data.json')))
OUT = os.path.join(HERE, 'source_data', 'SourceData.xlsx')


def sheet(wb, name, note, header, rows):
    ws = wb.create_sheet(name)
    ws.append([note]); ws['A1'].font = Font(italic=True)
    ws.append(header)
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([float(x) if isinstance(x, (np.floating, float)) else x for x in r])


if __name__ == '__main__':
    th, f3, f4, f5 = FD['theory'], FD['fig3'], FD['fig4'], FD['fig5']
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, 'Fig1c-d', 'Derived from Theorems 1 and 2 (closed form).',
          ['eta', 'C degree>=3', 'C chain', 'product weight w (chain)'],
          zip(th['eta'], [9.0] * len(th['eta']), th['C_chain'], th['w']))
    sheet(wb, 'Fig2a', 'Derived: two-sided bounds, sparse Pauli-Lindblad noise (9 generators, rate seed 4); '
          'pair depolarising exact constant in columns E-F.',
          ['average gate infidelity', 'lower bound', 'upper bound', '', 'pair-depolarising infidelity',
           'pair-depolarising constant'],
          [list(b) + ['', 0.75 * (1 - e), 9.0 if e * e <= 2 / 3 else 3 + 4 / (e * e)]
           for b, e in zip(th['bracket'] + [(None, None, None)] * 400, np.linspace(0.5, 0.998, 400))])
    sheet(wb, 'Fig2b', 'Derived: max{3/w, 18/(3-w)} at eta = 1.', ['w', 'worst-coordinate constant'],
          zip(th['ww'], th['Cw']))
    sheet(wb, 'Fig3a', 'Measured, ibm_cleveland job daoa3j8pqrnc7399miv0, product edge mode; simulator = noiseless '
          'Aer on the identical circuits.', ['source', 'realised sign', 'observed parity'],
          [('device', x, y) for x, y in f3['c110_edge_points_dev']] +
          [('simulator', x, y) for x, y in f3['c110_edge_points_sim']])
    rows = [('held-out run (daoa3j8pqrnc7399miv0)', k, v['n'], v['a_dev'], v['se_dev'], v['a_sim'], v['z'])
            for k, v in f3['c110'].items()]
    rows += [('target-resolved run (daobh65r85ps73ff233g)', k, v['n'], v['a'], v['se'], None, None)
             for k, v in f3['c112'].items()]
    sheet(wb, 'Fig3b', 'Fitted slope a of y = a x + c per mode; s.e. by bootstrap over circuits (400 resamples); '
          'plotted interval = 1.96 s.e.', ['run', 'mode', 'circuits', 'a', 's.e.', 'a (simulator)', 'z'], rows)
    sheet(wb, 'Fig4a', 'Fitted residual bias beta per target (sealed job daoduuo2fm4c73f5acpg); 95%% bootstrap '
          'interval, 1,000 joint resamples. tau = %g.' % FD['tau'],
          ['model', 'mode', 'beta', '2.5%', '97.5%'],
          [(m, md, f4[m]['per_mode'][md], f4[m]['per_mode_lo'][md], f4[m]['per_mode_hi'][md])
           for m in ('target', 'pooled') for md in f4[m]['per_mode']])
    sheet(wb, 'Fig4b', 'Bootstrap draws of the worst-coordinate bias B. UB95: target %.5f, pooled %.5f.'
          % (f4['target']['UB95'], f4['pooled']['UB95']), ['draw', 'B target-resolved', 'B pooled'],
          zip(range(1, len(f4['target']['Bboot']) + 1), f4['target']['Bboot'], f4['pooled']['Bboot']))
    sheet(wb, 'Fig5', 'Derived (model-based projection): risk ratio mixture/product; pointwise 95%% bootstrap bands; '
          'simultaneous one-sided 95%% bound on the worst-N ratio: target %.4f, pooled %.4f.'
          % (f5['target']['sim_hi'], f5['pooled']['sim_hi']),
          ['estimation copies N', 'target ratio', 'target 2.5%', 'target 97.5%', 'pooled ratio', 'pooled 2.5%',
           'pooled 97.5%'],
          zip(f5['N'], f5['target']['ratio'], f5['target']['ratio_lo'], f5['target']['ratio_hi'],
              f5['pooled']['ratio'], f5['pooled']['ratio_lo'], f5['pooled']['ratio_hi']))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print('sheets:', wb.sheetnames, '->', OUT)
